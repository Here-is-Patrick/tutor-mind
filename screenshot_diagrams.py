#!/usr/bin/env python3
"""Screenshot Mermaid diagrams from HTML and insert into Word report."""

import asyncio
from playwright.async_api import async_playwright
from docx import Document
from docx.shared import Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

HTML_PATH = "file:///workspace/diagrams.html"
REPORT_PATH = "/workspace/CS599_大作业报告.docx"
DIAGRAMS = [
    ("图1_系统整体架构图", 0),
    ("图2_Agent交互时序图", 1),
    ("图3_数据流图", 2),
    ("图4_记忆机制架构图", 3),
]


async def capture_diagrams():
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": 1600, "height": 1200})
        await page.goto(HTML_PATH, wait_until="networkidle")
        # Wait for mermaid to render
        await page.wait_for_timeout(3000)

        containers = await page.query_selector_all(".diagram-container")
        print(f"Found {len(containers)} diagram containers")

        for i, (name, idx) in enumerate(DIAGRAMS):
            if idx < len(containers):
                path = f"/workspace/{name}.png"
                await containers[idx].screenshot(path=path)
                print(f"Saved: {path}")

        await browser.close()


def insert_images_into_report():
    doc = Document(REPORT_PATH)

    # Map placeholder text to image file
    placeholders = {
        "【系统整体架构图": "图1_系统整体架构图.png",
        "【Agent 交互时序图": "图2_Agent交互时序图.png",
        "【数据流图": "图3_数据流图.png",
        "【记忆机制架构图": "图4_记忆机制架构图.png",
    }

    for paragraph in doc.paragraphs:
        for placeholder, img_file in placeholders.items():
            if placeholder in paragraph.text:
                # Clear placeholder and add image
                paragraph.clear()
                run = paragraph.add_run()
                run.add_picture(f"/workspace/{img_file}", width=Inches(5.8))
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                print(f"Inserted {img_file} into report")

    doc.save(REPORT_PATH)
    print(f"Report saved: {REPORT_PATH}")


if __name__ == "__main__":
    asyncio.run(capture_diagrams())
    insert_images_into_report()
